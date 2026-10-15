"""Report shaping and Excel export (BRD sections 4 and 5)."""
from __future__ import annotations

import io

import pandas as pd

from .logic import AMOUNT_DIFF_CODES, UNMATCHED_CODES, RunResult

EXCEL_MAX_ROWS = 1_048_575  # minus header


def _brd_columns(df: pd.DataFrame, label: str) -> pd.DataFrame:
    """Rename to the BRD output format: REFERENCE | NPA_STAGEID | PPG_DESC | Ops | Finance | Difference | Remarks."""
    ops, fin = f"{label} (Ops)", f"{label} (Finance)"
    rename = {
        "REFERENCE": "REFERENCE",
        "SCHDATE": "SCHDATE",
        "NPA_STAGEID": "NPA_STAGEID",
        "PPG_DESC": "PPG_DESC",
        "AMOUNT_OPS": ops,
        "AMOUNT_FINANCE": fin,
        "DIFFERENCE": f"Difference [{ops} - {fin}]",
        "REMARK_CODE": "Remark Code",
        "REMARKS": "Remarks",
        "REC_KEY": "FINREFERENCE_SCHDATE",
    }
    front = list(rename)
    rest = [c for c in df.columns if c not in front and c not in ("IN_OPS", "IN_FINANCE")]
    out = df[front + rest].rename(columns=rename)
    if "SCHDATE" in out:
        out["SCHDATE"] = pd.to_datetime(out["SCHDATE"], errors="coerce").dt.date
    return out


def primary_bounce(r: RunResult) -> pd.DataFrame:
    return _brd_columns(r.bounce, "Advise_Amount")


def primary_penal(r: RunResult) -> pd.DataFrame:
    return _brd_columns(r.penal, "Penal_Amount")


def _combined(r: RunResult, codes: set[int]) -> pd.DataFrame:
    keep = ["REC_KEY", "REFERENCE", "SCHDATE", "NPA_STAGEID", "PPG_DESC",
            "AMOUNT_OPS", "AMOUNT_FINANCE", "DIFFERENCE", "REMARK_CODE", "REMARKS",
            "LOAN_STATUS", "PRODUCTFLAG"]
    b = r.bounce[r.bounce["REMARK_CODE"].isin(codes)][keep].assign(CHARGE_TYPE="Bounce")
    p = r.penal[r.penal["REMARK_CODE"].isin(codes)][keep].assign(CHARGE_TYPE="Penal")
    df = pd.concat([b, p], ignore_index=True)
    out = _brd_columns(df, "Amount")
    out = out[["CHARGE_TYPE"] + [c for c in out.columns if c != "CHARGE_TYPE"]]
    return out.sort_values(["CHARGE_TYPE", "Remark Code", "REFERENCE"]).reset_index(drop=True)


def exceptions_amount_diff(r: RunResult) -> pd.DataFrame:
    return _combined(r, AMOUNT_DIFF_CODES)


def unmatched_records(r: RunResult) -> pd.DataFrame:
    return _combined(r, UNMATCHED_CODES)


def summary_table(r: RunResult) -> pd.DataFrame:
    s = r.summary.rename(columns={
        "CHARGE_TYPE": "Charge Type", "REMARK_CODE": "Remark Code", "REMARKS": "Remarks",
        "CASES": "Cases", "AMOUNT_OPS": "Amount (Ops)", "AMOUNT_FINANCE": "Amount (Finance)",
        "DIFFERENCE": "Difference [Ops - Finance]"})
    return s.sort_values("Remark Code").reset_index(drop=True)


REPORTS = {
    "primary": ("Primary Reconciliation Report", "Primary_Reconciliation_Report"),
    "exceptions": ("Exception Report – Amount Differences", "Exception_Report_Amount_Differences"),
    "unmatched": ("Unmatched Records Report", "Unmatched_Records_Report"),
}


def _report_sheets(r: RunResult, kind: str) -> tuple[dict[str, pd.DataFrame], set[int] | None]:
    b, p = primary_bounce(r), primary_penal(r)
    if kind == "primary":
        return {"Bounce_Recon": b, "Penal_Recon": p}, None
    codes = AMOUNT_DIFF_CODES if kind == "exceptions" else UNMATCHED_CODES
    return {
        "Bounce_" + ("Amount_Diff" if kind == "exceptions" else "Unmatched"): b[b["Remark Code"].isin(codes)],
        "Penal_" + ("Amount_Diff" if kind == "exceptions" else "Unmatched"): p[p["Remark Code"].isin(codes)],
    }, codes


def _write_workbook(r: RunResult, title_text: str, sheets: dict[str, pd.DataFrame],
                    summ: pd.DataFrame, info_extra: list[tuple[str, str]]) -> bytes:
    try:
        import xlsxwriter  # noqa: F401  (nicer formatting, faster)
        return _write_xlsxwriter(r, title_text, sheets, summ, info_extra)
    except ImportError:
        return _write_openpyxl(r, title_text, sheets, summ, info_extra)


def _split(df: pd.DataFrame):
    return [df.iloc[i:i + EXCEL_MAX_ROWS] for i in range(0, max(len(df), 1), EXCEL_MAX_ROWS)]


def _is_money(c: str) -> bool:
    return any(k in c for k in ("Amount", "Difference", "PENAL_"))


def _width(c: str) -> float:
    if c == "REFERENCE":
        return 24
    if c in ("FINREFERENCE_SCHDATE", "Remarks"):
        return 40
    return min(max(12, len(str(c)) * 0.9), 45)


def _write_xlsxwriter(r, title_text, sheets, summ, info_extra) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="xlsxwriter", datetime_format="dd-mmm-yyyy",
                        date_format="dd-mmm-yyyy") as xw:
        wb = xw.book
        hdr = wb.add_format({"bold": True, "bg_color": "#1F3864", "font_color": "white",
                             "text_wrap": True, "valign": "vcenter", "border": 1})
        money = wb.add_format({"num_format": "#,##0.00"})
        title = wb.add_format({"bold": True, "font_size": 14})
        bold = wb.add_format({"bold": True})

        ws = wb.add_worksheet("Summary")
        xw.sheets["Summary"] = ws
        ws.write(0, 0, f"Refund of Bounce and Penal for Excess Available Cases – {title_text}", title)
        info = [("Run ID", r.run_id), ("Run by", r.stats.get("user", ""))] + info_extra
        for i, (k, v) in enumerate(info, start=2):
            ws.write(i, 0, k, bold)
            ws.write(i, 1, v)
        start = len(info) + 4
        summ.to_excel(xw, sheet_name="Summary", startrow=start, index=False)
        for j, c in enumerate(summ.columns):
            ws.write(start, j, c, hdr)
        ws.set_column(0, 0, 34)
        ws.set_column(1, 2, 42)
        ws.set_column(3, 6, 18, money)

        for name, df in sheets.items():
            for n, part in enumerate(_split(df)):
                sn = name if n == 0 else f"{name[:27]}_{n + 1}"
                part.to_excel(xw, sheet_name=sn, index=False)
                w = xw.sheets[sn]
                for j, c in enumerate(part.columns):
                    w.write(0, j, c, hdr)
                    w.set_column(j, j, _width(c), money if _is_money(c) else None)
                w.set_row(0, 45)
                w.freeze_panes(1, 0)
                w.autofilter(0, 0, max(len(part), 1), len(part.columns) - 1)
    return buf.getvalue()


def _write_openpyxl(r, title_text, sheets, summ, info_extra) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hdr_font, hdr_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="DD-MMM-YYYY",
                        date_format="DD-MMM-YYYY") as xw:
        info = [("Run ID", r.run_id), ("Run by", r.stats.get("user", ""))] + info_extra
        start = len(info) + 4
        summ.to_excel(xw, sheet_name="Summary", startrow=start, index=False)
        ws = xw.sheets["Summary"]
        ws.cell(1, 1, f"Refund of Bounce and Penal for Excess Available Cases – {title_text}").font = Font(bold=True, size=14)
        for i, (k, v) in enumerate(info, start=3):
            ws.cell(i, 1, k).font = Font(bold=True)
            ws.cell(i, 2, v)
        for j in range(1, len(summ.columns) + 1):
            c = ws.cell(start + 1, j)
            c.font, c.fill = hdr_font, hdr_fill
        for col, wdt in (("A", 34), ("B", 42), ("C", 42), ("D", 18), ("E", 18), ("F", 18), ("G", 18)):
            ws.column_dimensions[col].width = wdt

        for name, df in sheets.items():
            for n, part in enumerate(_split(df)):
                sn = name if n == 0 else f"{name[:27]}_{n + 1}"
                part.to_excel(xw, sheet_name=sn, index=False)
                w = xw.sheets[sn]
                for j, c in enumerate(part.columns, start=1):
                    cell = w.cell(1, j)
                    cell.font, cell.fill = hdr_font, hdr_fill
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    letter = get_column_letter(j)
                    w.column_dimensions[letter].width = _width(c)
                    if _is_money(c):
                        for row in w.iter_rows(min_row=2, min_col=j, max_col=j):
                            row[0].number_format = "#,##0.00"
                w.freeze_panes = "A2"
                w.auto_filter.ref = w.dimensions
    return buf.getvalue()


def build_reports(r: RunResult, ops_name: str, fin_name: str, settings_desc: dict) -> dict[str, tuple[str, str, bytes]]:
    """Return {kind: (label, file_name, xlsx_bytes)} for the three BRD deliverables (section 5)."""
    base_info = [
        ("OPS file", ops_name), ("Finance file", fin_name),
        ("OPS rows / LAN+SCH keys", f"{r.stats['ops_rows']:,} / {r.stats['ops_keys']:,}"),
        ("Finance rows / LAN+SCH keys", f"{r.stats['fin_rows']:,} / {r.stats['fin_keys']:,}"),
        ("Keys repeated within OPS / Finance",
         f"{r.stats['ops_repeated_keys']:,} / {r.stats['fin_repeated_keys']:,}"),
    ] + [(k, str(v)) for k, v in settings_desc.items()]
    full_summary = summary_table(r)
    out = {}
    for kind, (label, stem) in REPORTS.items():
        sheets, codes = _report_sheets(r, kind)
        summ = full_summary if codes is None else full_summary[full_summary["Remark Code"].isin(codes)]
        info = base_info + ([("Tables", ", ".join(v for k, v in r.tables.items()))] if kind == "primary" else [])
        out[kind] = (label, f"{stem}_{r.run_id}.xlsx",
                     _write_workbook(r, label, sheets, summ.reset_index(drop=True), info))
    return out
